"""
server.py – FastAPI REST API Backend for LearnMate.

Preserves 100% of all existing Python services, ML models, database models,
and Gemini AI pipelines without altering any underlying business logic.
"""
from __future__ import annotations

import io
import json
from collections import defaultdict
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from pydantic import BaseModel, Field

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))

import config
from db.session import init_db, get_db
from db.models import User, Subject, Topic, Session, Quiz, QuizAttempt, FlashcardSet, Flashcard, Doubt
from services.mastery import confidence_to_theta, theta_to_confidence, update_theta
from services.planner import build_plan, should_replan, schedule_count_for_subject
from services.predictor import (
    predict,
    load_model,
    warmup,
    session_suggestion,
    skip_impact,
    build_user_readiness,
    to_percent,
    visualiser_overview,
    visualiser_state,
)
from services.ml_analyzer import analyze_user_weaknesses, apply_dynamic_timetable_update, add_subject_with_topics
import services.llm as llm_service
import services.doubt_solver as doubt_solver

# Ensure DB is initialized
init_db()


@asynccontextmanager
async def lifespan(_app: FastAPI):
    warmup()
    yield


app = FastAPI(title="LearnMate API", version="2.0", lifespan=lifespan)

# Enable CORS for local Vite dev server
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def _get_or_create_user(db):
    user = db.query(User).first()
    if not user:
        user = User(
            name="Alex",
            timezone="UTC",
            study_days="Mon,Tue,Wed,Thu,Fri,Sat",
            hours_per_day=3.0,
            session_min=config.DEFAULT_SESSION_MINUTES,
            break_min=config.DEFAULT_BREAK_MINUTES,
            prev_exam_score=68.0,
            weekly_study_hours=8.0
        )
        db.add(user)
        db.commit()
        db.refresh(user)
    return user


def _sessions_by_subject(db, user_id: int) -> dict:
    grouped = defaultdict(list)
    for sess in db.query(Session).filter(Session.user_id == user_id).all():
        grouped[sess.subject_id].append(sess)
    return grouped


def _readiness_payload(db, user) -> dict:
    subjects = db.query(Subject).filter(Subject.user_id == user.id).all()
    grouped = _sessions_by_subject(db, user.id)
    raw = build_user_readiness(
        subjects,
        grouped,
        user.session_min or config.DEFAULT_SESSION_MINUTES,
        fallback_previous_score=user.prev_exam_score,
    )
    subjects_out = []
    for row in raw.get("subjects", []):
        subjects_out.append({
            **row,
            "pass_probability": to_percent(row.get("pass_probability")),
        })

    pass_threshold = round(config.VISUALISER_DECISION_THRESHOLD * 100, 1)
    scored = [row for row in subjects_out if row.get("pass_probability") is not None]
    limiting = min(scored, key=lambda row: row["pass_probability"]) if scored else None
    above = sum(1 for row in scored if row["pass_probability"] >= pass_threshold)

    if not subjects_out:
        note = "Add subjects with a previous score to estimate readiness."
    elif not scored:
        note = "Add a previous score to your subjects to estimate readiness."
    elif limiting and limiting["pass_probability"] < pass_threshold:
        name = limiting.get("subject_name") or "A subject"
        pct = round(limiting["pass_probability"], 1)
        scheduled = float(limiting.get("scheduled_hours") or 0.0)
        benchmark = limiting.get("benchmark_hours")
        score = limiting.get("previous_score")
        if benchmark is not None and scheduled + 0.05 < float(benchmark):
            need = max(0.5, round(float(benchmark) - scheduled, 1))
            note = (
                f"{name} is the bottleneck at {pct}% — add {need} more planned "
                f"hours (currently {scheduled:g}/{float(benchmark):g} h)."
            )
        elif score is not None and float(score) < config.PASS_TARGET_PERCENTAGE:
            note = (
                f"{name} is the bottleneck at {pct}% because its previous score "
                f"({round(float(score))}) is below the dataset's "
                f"{round(config.PASS_TARGET_PERCENTAGE)}% pass mark."
            )
        else:
            note = (
                f"{name} is the bottleneck at {pct}%; add study time or run "
                "Smart Optimize to lift it."
            )
    elif above == len(scored):
        note = (
            f"All {len(scored)} subjects are above the {round(pass_threshold)}% "
            "pass threshold — keep the plan balanced."
        )
    else:
        note = (
            f"{above}/{len(scored)} subjects are above the "
            f"{round(pass_threshold)}% pass threshold; "
            f"{limiting['subject_name']} is lowest at "
            f"{round(limiting['pass_probability'])}%."
        )

    return {
        "model_available": raw.get("model_available", False),
        "overall_pass_probability": to_percent(raw.get("overall_pass_probability")),
        "subjects": subjects_out,
        "pass_threshold": pass_threshold,
        "subjects_scored": len(scored),
        "subjects_above_pass": above,
        "limiting_subject": (
            {
                "subject_id": limiting.get("subject_id"),
                "subject_name": limiting.get("subject_name"),
                "pass_probability": limiting.get("pass_probability"),
                "scheduled_hours": limiting.get("scheduled_hours"),
                "benchmark_hours": limiting.get("benchmark_hours"),
                "previous_score": limiting.get("previous_score"),
            }
            if limiting
            else None
        ),
        "note": note,
        "warning": None if raw.get("model_available") else "Pass estimates unavailable — the trained model did not load.",
    }


# ─────────────────────────────────────────────────────────────────────────────
# 1. System Status & Diagnostics
# ─────────────────────────────────────────────────────────────────────────────

@app.get("/api/status")
def get_system_status():
    model = load_model()
    gemini = llm_service.gemini_status()
    return {
        "ml_model_active": model is not None,
        "ml_model_name": "Random Forest Classifier" if model is not None else "Not Found",
        # gemini_active stays true when a key is present (backwards compatible).
        # gemini_verified only becomes true after a successful Gemini call.
        "gemini_active": gemini["configured"],
        "gemini_configured": gemini["configured"],
        "gemini_verified": bool(gemini["active_model"]),
        "gemini_model": gemini["preferred_model"],
        "gemini_active_model": gemini["active_model"],
        "gemini_fallback_models": list(getattr(config, "GEMINI_FALLBACK_MODELS", [])),
        "gemini_last_error": gemini["last_error"],
        "groq_active": config.GROQ_ENABLED,
        "groq_model": config.GROQ_MODEL,
        "api_status": "online",
        "version": "2.0-react",
    }


# ─────────────────────────────────────────────────────────────────────────────
# 2. User & Dashboard
# ─────────────────────────────────────────────────────────────────────────────

@app.get("/api/user")
def get_user_profile():
    db = get_db()
    try:
        user = _get_or_create_user(db)
        return {
            "id": user.id,
            "name": user.name,
            "study_days": [d.strip() for d in user.study_days.split(",") if d.strip()],
            "hours_per_day": user.hours_per_day,
            "session_min": user.session_min,
            "break_min": user.break_min,
            "prev_exam_score": user.prev_exam_score,
            "weekly_study_hours": user.weekly_study_hours,
        }
    finally:
        db.close()


class UserUpdateModel(BaseModel):
    name: Optional[str] = None
    study_days: Optional[list[str]] = None
    hours_per_day: Optional[float] = None
    session_min: Optional[int] = None
    break_min: Optional[int] = None
    prev_exam_score: Optional[float] = None
    weekly_study_hours: Optional[float] = None


@app.put("/api/user")
def update_user_profile(body: UserUpdateModel):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        if body.name is not None:
            user.name = body.name.strip()
        if body.study_days is not None:
            user.study_days = ",".join(body.study_days)
        if body.hours_per_day is not None:
            user.hours_per_day = body.hours_per_day
        if body.session_min is not None:
            user.session_min = body.session_min
        if body.break_min is not None:
            user.break_min = body.break_min
        if body.prev_exam_score is not None:
            user.prev_exam_score = body.prev_exam_score
        if body.weekly_study_hours is not None:
            user.weekly_study_hours = body.weekly_study_hours
        db.commit()
        return {"ok": True, "message": "Profile updated."}
    finally:
        db.close()


@app.get("/api/dashboard")
def get_dashboard():
    db = get_db()
    try:
        user = _get_or_create_user(db)
        subjects = db.query(Subject).filter(Subject.user_id == user.id).all()
        topics = db.query(Topic).filter(Topic.user_id == user.id).all()
        all_sessions = db.query(Session).filter(Session.user_id == user.id).all()

        today_str = date.today().isoformat()
        today_sessions = (
            db.query(Session)
            .filter(
                Session.user_id == user.id,
                Session.scheduled_start >= today_str,
                Session.scheduled_start < today_str + "T23:59:59",
            )
            .order_by(Session.scheduled_start)
            .all()
        )

        completed_sessions = sum(1 for s in all_sessions if s.status == "completed")

        readiness = _readiness_payload(db, user)

        # Run ML weakness analysis
        ml_res = analyze_user_weaknesses(db, user.id)
        critical_topics = ml_res.get("critical_topics", [])[:4]

        # Upcoming sessions list
        sessions_data = []
        for s in today_sessions:
            sub = db.get(Subject, s.subject_id)
            top = db.get(Topic, s.topic_id) if s.topic_id else None
            sessions_data.append({
                "id": s.id,
                "subject_id": s.subject_id,
                "scheduled_start": s.scheduled_start,
                "scheduled_end": s.scheduled_end,
                "time": s.scheduled_start[11:16],
                "start_time": s.scheduled_start[11:16],
                "end_time": s.scheduled_end[11:16] if s.scheduled_end else None,
                "topic_name": top.name if top else "General Study",
                "subject_name": sub.name if sub else "—",
                "duration_min": s.duration_min,
                "status": s.status,
                "notes": s.notes or "",
            })

        # Calculate average mastery
        avg_mastery = 0.0
        if topics:
            confs = [theta_to_confidence(t.theta or 0.0) * 100 for t in topics]
            avg_mastery = round(sum(confs) / len(confs), 1)

        # Real syllabus coverage: how many distinct topics appear in the plan.
        topics_in_plan = len({s.topic_id for s in all_sessions if s.topic_id})
        syllabus_progress = round((topics_in_plan / len(topics)) * 100, 1) if topics else 0.0
        planned_minutes = sum(
            s.duration_min or 0 for s in all_sessions if s.status in ("scheduled", "completed")
        )

        # Next future slot, so an evening visit still shows what comes next
        # instead of pretending the timetable is empty.
        now_iso = datetime.now(timezone.utc).isoformat()
        next_sess = (
            db.query(Session)
            .filter(
                Session.user_id == user.id,
                Session.status == "scheduled",
                Session.scheduled_start >= now_iso,
            )
            .order_by(Session.scheduled_start)
            .first()
        )
        next_session = None
        if next_sess:
            next_subject = db.get(Subject, next_sess.subject_id)
            next_topic = db.get(Topic, next_sess.topic_id) if next_sess.topic_id else None
            next_session = {
                "id": next_sess.id,
                "date": next_sess.scheduled_start[:10],
                "start_time": next_sess.scheduled_start[11:16],
                "end_time": next_sess.scheduled_end[11:16] if next_sess.scheduled_end else None,
                "subject_name": next_subject.name if next_subject else "—",
                "topic_name": next_topic.name if next_topic else "General Study",
            }

        return {
            "user_name": user.name,
            "user": {
                "name": user.name,
                "weekly_study_hours": user.weekly_study_hours,
                "prev_exam_score": user.prev_exam_score,
                "session_min": user.session_min,
                "break_min": user.break_min,
                "hours_per_day": user.hours_per_day,
                "study_days": [d.strip() for d in user.study_days.split(",") if d.strip()],
            },
            "pass_probability": readiness.get("overall_pass_probability"),
            "readiness": readiness,
            "thresholds": {
                "pass": round(config.VISUALISER_DECISION_THRESHOLD * 100, 1),
                "on_track": config.PASS_TARGET_PERCENTAGE,
            },
            "total_subjects": len(subjects),
            "total_topics": len(topics),
            "topics_in_plan": topics_in_plan,
            "syllabus_progress": syllabus_progress,
            "planned_hours": round(planned_minutes / 60.0, 1),
            "completed_sessions": completed_sessions,
            "average_mastery": avg_mastery,
            "today_sessions": sessions_data,
            "next_session": next_session,
            "critical_topics": critical_topics,
            "subject_risk": ml_res.get("subject_risk", []),
        }
    finally:
        db.close()


# ─────────────────────────────────────────────────────────────────────────────
# 3. Subjects & Topics
# ─────────────────────────────────────────────────────────────────────────────

@app.get("/api/subjects")
def get_subjects():
    db = get_db()
    try:
        user = _get_or_create_user(db)
        subjects = db.query(Subject).filter(Subject.user_id == user.id).all()
        res = []
        for s in subjects:
            topics_data = []
            for t in s.topics:
                conf = round(theta_to_confidence(t.theta or 0.0) * 100, 1)
                topics_data.append({
                    "id": t.id,
                    "name": t.name,
                    "difficulty": t.difficulty,
                    "confidence": conf,
                    "theta": round(t.theta or 0.0, 3),
                    "n_answered": t.n_answered or 0,
                    "estimated_min": t.estimated_min or 60,
                    "status": t.status,
                })
            res.append({
                "id": s.id,
                "name": s.name,
                "exam_date": s.exam_date,
                "priority": s.priority,
                "previous_score": s.previous_score,
                "topics": topics_data,
            })
        return res
    finally:
        db.close()


class CreateSubjectModel(BaseModel):
    name: str
    exam_date: str
    priority: int = 1
    previous_score: float = Field(..., ge=0, le=100)
    topics: list[str]
    accept_suggested_sessions: bool = False
    suggested_session_count: Optional[int] = None


@app.post("/api/subjects")
def create_subject(body: CreateSubjectModel):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        if not body.name.strip():
            raise HTTPException(status_code=400, detail="Subject name is required.")
        if not body.topics:
            raise HTTPException(status_code=400, detail="At least one topic is required.")

        subj = Subject(
            user_id=user.id,
            name=body.name.strip(),
            exam_date=body.exam_date,
            priority=body.priority,
            previous_score=body.previous_score,
        )
        db.add(subj)
        db.flush()

        for tname in body.topics:
            tname = tname.strip()
            if tname:
                db.add(Topic(
                    user_id=user.id,
                    subject_id=subj.id,
                    name=tname,
                    difficulty="medium",
                    confidence=50.0,
                    theta=confidence_to_theta(50.0),
                    estimated_min=60,
                ))
        scheduled = 0
        if body.accept_suggested_sessions:
            n = body.suggested_session_count
            if n is None:
                sug = session_suggestion(body.previous_score, user.session_min or config.DEFAULT_SESSION_MINUTES)
                n = sug.get("suggested_session_count") or 0
            scheduled = schedule_count_for_subject(db, user.id, subj, int(n))
        db.commit()
        return {
            "ok": True,
            "subject_id": subj.id,
            "sessions_created": scheduled,
            "message": "Subject created successfully.",
        }
    finally:
        db.close()


class UpdateSubjectModel(BaseModel):
    name: Optional[str] = None
    exam_date: Optional[str] = None
    priority: Optional[int] = None
    previous_score: Optional[float] = Field(None, ge=0, le=100)


@app.put("/api/subjects/{subject_id}")
def update_subject(subject_id: int, body: UpdateSubjectModel):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        subj = db.query(Subject).filter(Subject.id == subject_id, Subject.user_id == user.id).first()
        if not subj:
            raise HTTPException(status_code=404, detail="Subject not found.")
        if body.name is not None:
            subj.name = body.name.strip()
        if body.exam_date is not None:
            subj.exam_date = body.exam_date
        if body.priority is not None:
            subj.priority = body.priority
        if body.previous_score is not None:
            subj.previous_score = body.previous_score
        db.commit()
        return {"ok": True, "message": "Subject updated."}
    finally:
        db.close()


@app.delete("/api/subjects/{subject_id}")
def delete_subject(subject_id: int):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        subj = db.query(Subject).filter(Subject.id == subject_id, Subject.user_id == user.id).first()
        if not subj:
            raise HTTPException(status_code=404, detail="Subject not found.")
        db.delete(subj)
        db.commit()
        return {"ok": True, "message": "Subject deleted."}
    finally:
        db.close()


class AddTopicModel(BaseModel):
    name: str
    difficulty: str = "medium"
    confidence: float = 50.0
    estimated_min: int = 60


@app.post("/api/subjects/{subject_id}/topics")
def add_topic(subject_id: int, body: AddTopicModel):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        subj = db.query(Subject).filter(Subject.id == subject_id, Subject.user_id == user.id).first()
        if not subj:
            raise HTTPException(status_code=404, detail="Subject not found.")
        topic = Topic(
            user_id=user.id,
            subject_id=subj.id,
            name=body.name.strip(),
            difficulty=body.difficulty,
            confidence=body.confidence,
            theta=confidence_to_theta(body.confidence),
            estimated_min=body.estimated_min,
        )
        db.add(topic)
        db.commit()
        return {"ok": True, "topic_id": topic.id, "message": "Topic added."}
    finally:
        db.close()


class UpdateTopicModel(BaseModel):
    difficulty: Optional[str] = None
    confidence: Optional[float] = None
    estimated_min: Optional[int] = None


@app.put("/api/topics/{topic_id}")
def update_topic(topic_id: int, body: UpdateTopicModel):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        topic = db.query(Topic).filter(Topic.id == topic_id, Topic.user_id == user.id).first()
        if not topic:
            raise HTTPException(status_code=404, detail="Topic not found.")
        if body.difficulty is not None:
            topic.difficulty = body.difficulty
        if body.confidence is not None:
            topic.confidence = body.confidence
            topic.theta = confidence_to_theta(body.confidence)
        if body.estimated_min is not None:
            topic.estimated_min = body.estimated_min
        db.commit()
        return {"ok": True, "message": "Topic updated."}
    finally:
        db.close()


@app.delete("/api/topics/{topic_id}")
def delete_topic(topic_id: int):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        topic = db.query(Topic).filter(Topic.id == topic_id, Topic.user_id == user.id).first()
        if not topic:
            raise HTTPException(status_code=404, detail="Topic not found.")
        db.delete(topic)
        db.commit()
        return {"ok": True, "message": "Topic deleted."}
    finally:
        db.close()


# ─────────────────────────────────────────────────────────────────────────────
# 4. Study Plan / Timetable
# ─────────────────────────────────────────────────────────────────────────────

@app.get("/api/study-plan")
def get_study_plan(date_str: Optional[str] = None):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        query = db.query(Session).filter(Session.user_id == user.id)
        if date_str:
            query = query.filter(
                Session.scheduled_start >= date_str,
                Session.scheduled_start < date_str + "T23:59:59"
            )
        sessions = query.order_by(Session.scheduled_start).all()

        res = []
        for s in sessions:
            subj = db.get(Subject, s.subject_id)
            top = db.get(Topic, s.topic_id) if s.topic_id else None
            res.append({
                "id": s.id,
                "subject_id": s.subject_id,
                "scheduled_start": s.scheduled_start,
                "scheduled_end": s.scheduled_end,
                "time": s.scheduled_start[11:16],
                "end_time": s.scheduled_end[11:16] if s.scheduled_end else None,
                "date": s.scheduled_start[:10],
                "topic_name": top.name if top else "General Study",
                "subject_name": subj.name if subj else "—",
                "duration_min": s.duration_min,
                "status": s.status,
                "notes": s.notes or "",
                "is_weak_priority": bool(s.notes and ("ML Priority Focus" in s.notes or "Weak Topic" in s.notes)),
            })

        needs_replan, reason = should_replan(db, user.id)
        return {
            "sessions": res,
            "needs_replan": needs_replan,
            "replan_reason": reason,
            "readiness": _readiness_payload(db, user),
        }
    finally:
        db.close()


@app.post("/api/study-plan/generate")
def generate_standard_plan():
    db = get_db()
    try:
        user = _get_or_create_user(db)
        res = build_plan(db, user.id)
        res["readiness"] = _readiness_payload(db, user)
        return res
    finally:
        db.close()


class MLOptimizeModel(BaseModel):
    custom_weak_topics: Optional[list[str]] = None


@app.post("/api/study-plan/ml-optimize")
def optimize_plan_ml(body: MLOptimizeModel = None):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        before = _readiness_payload(db, user).get("overall_pass_probability")
        weak_names = body.custom_weak_topics if body else None
        res = apply_dynamic_timetable_update(db, user.id, custom_weak_topic_names=weak_names)
        after = _readiness_payload(db, user).get("overall_pass_probability")
        res["readiness_before"] = before
        res["readiness_after"] = after
        res["readiness_delta"] = (
            round(after - before, 1)
            if before is not None and after is not None
            else None
        )
        return res
    finally:
        db.close()


class SessionActionModel(BaseModel):
    status: Optional[str] = None
    shift_days: Optional[int] = None


@app.put("/api/sessions/{session_id}")
def update_session(session_id: int, body: SessionActionModel):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        sess = db.query(Session).filter(Session.id == session_id, Session.user_id == user.id).first()
        if not sess:
            raise HTTPException(status_code=404, detail="Session not found.")

        if body.status:
            sess.status = body.status
            if body.status == "completed":
                sess.actual_end = datetime.now(timezone.utc).isoformat()
        if body.shift_days:
            try:
                s = datetime.fromisoformat(sess.scheduled_start)
                e = datetime.fromisoformat(sess.scheduled_end)
                sess.scheduled_start = (s + timedelta(days=body.shift_days)).isoformat()
                sess.scheduled_end = (e + timedelta(days=body.shift_days)).isoformat()
            except Exception:
                pass
        db.commit()
        return {"ok": True, "message": "Session updated.", "readiness": _readiness_payload(db, user)}
    finally:
        db.close()


@app.get("/api/sessions/{session_id}/skip-preview")
def preview_session_skip(session_id: int):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        sess = db.query(Session).filter(Session.id == session_id, Session.user_id == user.id).first()
        if not sess:
            raise HTTPException(status_code=404, detail="Session not found.")
        subj = db.get(Subject, sess.subject_id)
        if not subj:
            raise HTTPException(status_code=404, detail="Subject not found.")
        subject_sessions = (
            db.query(Session)
            .filter(Session.user_id == user.id, Session.subject_id == subj.id)
            .all()
        )
        impact = skip_impact(subj, subject_sessions, sess, user.session_min or config.DEFAULT_SESSION_MINUTES)
        return {
            "model_available": impact["model_available"],
            "subject_name": impact["subject_name"],
            "before": to_percent(impact["before"]),
            "after": to_percent(impact["after"]),
            "drop_percentage_points": impact["drop_percentage_points"],
            "warn": impact["warn"],
            "scheduled_hours_before": impact["scheduled_hours_before"],
            "scheduled_hours_after": impact["scheduled_hours_after"],
            "warning": None if impact["model_available"] else "Pass estimates unavailable — the trained model did not load.",
        }
    finally:
        db.close()


@app.delete("/api/sessions/{session_id}")
def delete_session(session_id: int):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        sess = db.query(Session).filter(Session.id == session_id, Session.user_id == user.id).first()
        if not sess:
            raise HTTPException(status_code=404, detail="Session not found.")
        db.delete(sess)
        db.commit()
        return {"ok": True, "message": "Session deleted.", "readiness": _readiness_payload(db, user)}
    finally:
        db.close()


@app.get("/api/ml-diagnostic")
def get_ml_diagnostic():
    db = get_db()
    try:
        user = _get_or_create_user(db)
        res = analyze_user_weaknesses(db, user.id)
        if res.get("ok"):
            res["weak_topics"] = [
                {
                    "subject": t.get("subject_name"),
                    "topic": t.get("name"),
                    "weakness_score": t.get("weakness_score"),
                    "confidence": t.get("confidence"),
                    "recommendation": t.get("recommendation"),
                }
                for t in res.get("critical_topics", [])
            ]
            scheduled = (
                db.query(Session)
                .filter(
                    Session.user_id == user.id,
                    Session.status == "scheduled",
                )
                .all()
            )
            priority_slots = sum(
                1
                for s in scheduled
                if s.notes and ("ML Priority Focus" in s.notes or "Weak Topic" in s.notes)
            )
            res["priority_slots"] = priority_slots
            res["scheduled_sessions"] = len(scheduled)
        return res
    finally:
        db.close()


# ─────────────────────────────────────────────────────────────────────────────
# 5. AI Quizzes
# ─────────────────────────────────────────────────────────────────────────────

class QuizGenerateModel(BaseModel):
    subject: str
    topic: str
    difficulty: str = "medium"
    n_questions: int = 5


@app.post("/api/quizzes/generate")
def generate_ai_quiz(body: QuizGenerateModel):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        res = llm_service.generate_quiz(
            subject=body.subject,
            topic=body.topic,
            difficulty=body.difficulty,
            n_questions=body.n_questions,
            db_session=db
        )
        if not res["ok"]:
            raise HTTPException(status_code=400, detail=res.get("error", "Failed to generate quiz."))

        # Cache quiz in DB
        sub = db.query(Subject).filter(Subject.user_id == user.id, Subject.name == body.subject).first()
        top = db.query(Topic).filter(Topic.user_id == user.id, Topic.name == body.topic).first()

        quiz = Quiz(
            user_id=user.id,
            subject_id=sub.id if sub else None,
            topic_id=top.id if top else None,
            difficulty=body.difficulty,
            n_questions=body.n_questions,
            questions=json.dumps(res["questions"]),
            cache_key=llm_service._make_cache_key(body.subject, body.topic, body.difficulty, body.n_questions)
        )
        db.add(quiz)
        db.commit()

        return {
            "ok": True,
            "quiz_id": quiz.id,
            "topic_id": top.id if top else None,
            "questions": res["questions"],
            "provider": res.get("provider"),
            "provider_error": res.get("provider_error"),
        }
    finally:
        db.close()


class QuizSubmitModel(BaseModel):
    quiz_id: Optional[int] = 0
    topic_id: Optional[int] = None
    questions: Optional[list[dict]] = None
    answers: list[object]


@app.post("/api/quizzes/submit")
def submit_quiz_answers(body: QuizSubmitModel):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        questions = body.questions or []
        if not questions and body.quiz_id:
            saved_q = db.get(Quiz, body.quiz_id)
            if saved_q and saved_q.questions:
                try:
                    questions = json.loads(saved_q.questions)
                except Exception:
                    pass

        # Normalize answers to integer list
        normalized_answers = []
        for a in body.answers:
            if isinstance(a, dict):
                normalized_answers.append(int(a.get("selected_index", -1)))
            elif isinstance(a, int):
                normalized_answers.append(a)
            else:
                try:
                    normalized_answers.append(int(a))
                except Exception:
                    normalized_answers.append(-1)

        correct = sum(1 for i, q in enumerate(questions) if i < len(normalized_answers) and normalized_answers[i] == q.get("correct_index"))
        score_frac = correct / len(questions) if questions else 0.0

        # Update Bayesian IRT mastery for topic
        if body.topic_id:
            from services.mastery import apply_quiz_result
            apply_quiz_result(db, body.topic_id, normalized_answers, questions)

        # Categorize concepts and track WRONG QUESTIONS
        weak_topics = []
        mastered_topics = []
        wrong_questions = []

        for i, q in enumerate(questions):
            c_tag = q.get("topic") or "General Concept"
            sel_idx = normalized_answers[i] if i < len(normalized_answers) else -1
            correct_idx = q.get("correct_index", 0)
            options = q.get("options", [])
            is_correct = (sel_idx == correct_idx)

            if is_correct:
                mastered_topics.append(c_tag)
            else:
                weak_topics.append(c_tag)
                user_opt_text = options[sel_idx] if 0 <= sel_idx < len(options) else "Unanswered"
                correct_opt_text = options[correct_idx] if 0 <= correct_idx < len(options) else ""
                wrong_questions.append({
                    "question_index": i + 1,
                    "question": q.get("question", ""),
                    "topic": c_tag,
                    "user_answer": user_opt_text,
                    "correct_answer": correct_opt_text,
                    "explanation": q.get("explanation", ""),
                    "difficulty": q.get("difficulty", "medium"),
                })

        unique_weak = list(dict.fromkeys(weak_topics))
        unique_mastered = [m for m in dict.fromkeys(mastered_topics) if m not in unique_weak]

        # Trigger dynamic ML timetable update if weak topics found
        replan_triggered = False
        replan_msg = ""
        if unique_weak:
            try:
                plan_res = apply_dynamic_timetable_update(db, user.id, custom_weak_topics=unique_weak)
                replan_triggered = True
                replan_msg = f"Timetable automatically recalibrated: {plan_res.get('weak_sessions_scheduled', 0)} reinforcement sessions added for {', '.join(unique_weak[:2])}."
            except Exception:
                pass

        # Save QuizAttempt with guaranteed valid foreign key
        target_quiz_id = body.quiz_id
        if not target_quiz_id or target_quiz_id == 0:
            existing_quiz = db.query(Quiz).filter(Quiz.user_id == user.id).first()
            if existing_quiz:
                target_quiz_id = existing_quiz.id
            else:
                new_q = Quiz(
                    user_id=user.id,
                    subject_id=None,
                    difficulty="medium",
                    n_questions=len(questions),
                    questions=json.dumps(questions),
                )
                db.add(new_q)
                db.flush()
                target_quiz_id = new_q.id

        qa = QuizAttempt(
            user_id=user.id,
            quiz_id=target_quiz_id,
            answers=json.dumps(normalized_answers),
            score=score_frac,
        )
        db.add(qa)
        db.commit()

        return {
            "score": correct,
            "total": len(questions),
            "score_pct": round(score_frac * 100, 1),
            "score_percentage": round(score_frac * 100, 1),
            "weak_topics": unique_weak,
            "mastered_topics": unique_mastered,
            "wrong_questions": wrong_questions,
            "replan_triggered": replan_triggered,
            "replan_message": replan_msg,
        }
    finally:
        db.close()


@app.get("/api/quizzes/history")
def get_quiz_history():
    db = get_db()
    try:
        user = _get_or_create_user(db)
        attempts = (
            db.query(QuizAttempt)
            .filter(QuizAttempt.user_id == user.id)
            .order_by(QuizAttempt.submitted_at.desc())
            .limit(10)
            .all()
        )
        res = []
        for a in attempts:
            q = db.get(Quiz, a.quiz_id)
            res.append({
                "id": a.id,
                "score_pct": int(a.score * 100),
                "n_questions": q.n_questions if q else "?",
                "difficulty": q.difficulty if q else "medium",
                "date": a.submitted_at.strftime("%d %b, %H:%M") if a.submitted_at else "—",
            })
        return res
    finally:
        db.close()


# ─────────────────────────────────────────────────────────────────────────────
# 6. Document / PPT Upload & Flashcards
# ─────────────────────────────────────────────────────────────────────────────

def _extract_shape_text(shape) -> list[str]:
    """Recursively extract text from pptx shapes including tables and groups."""
    texts = []
    try:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                t = paragraph.text.strip()
                if t:
                    texts.append(t)
        if shape.has_table:
            for row in shape.table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    texts.append(" | ".join(row_texts))
        if hasattr(shape, "shapes"):
            for sub_shape in shape.shapes:
                texts.extend(_extract_shape_text(sub_shape))
    except Exception:
        pass
    return texts


def _extract_ppt_fallback(raw_bytes: bytes) -> Optional[str]:
    """Extract readable text strings from legacy binary .ppt files."""
    import re
    chunks = []
    try:
        decoded_utf16 = raw_bytes.decode("utf-16le", errors="ignore")
        matches = re.findall(r"[\w\s.,;:?!()/\-]{6,}", decoded_utf16)
        chunks.extend([m.strip() for m in matches if len(m.strip()) > 8])
    except Exception:
        pass
    try:
        decoded_ascii = raw_bytes.decode("latin1", errors="ignore")
        matches = re.findall(r"[A-Za-z0-9\s.,;:?!()/\-]{6,}", decoded_ascii)
        chunks.extend([m.strip() for m in matches if len(m.strip()) > 8])
    except Exception:
        pass

    seen = set()
    cleaned = []
    for c in chunks:
        if c not in seen and not c.startswith("Microsoft") and not c.startswith("PowerPoint"):
            seen.add(c)
            cleaned.append(c)
    return "\n\n".join(cleaned[:300]) if cleaned else None


def _extract_text(filename: str, raw_bytes: bytes) -> Optional[str]:
    fname = filename.lower()
    if fname.endswith(".txt"):
        return raw_bytes.decode("utf-8", errors="replace")

    if fname.endswith(".pdf"):
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw_bytes))
            pages = [p.extract_text() or "" for p in reader.pages]
            text = "\n".join(pages).strip()
            return text if text else None
        except Exception:
            return None

    if fname.endswith(".docx"):
        try:
            import docx
            doc = docx.Document(io.BytesIO(raw_bytes))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip()).strip()
            return text if text else None
        except Exception:
            return None

    if fname.endswith(".pptx") or fname.endswith(".ppt"):
        try:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(raw_bytes))
            text_runs = []
            for slide_num, slide in enumerate(prs.slides, 1):
                slide_texts = []
                for shape in slide.shapes:
                    slide_texts.extend(_extract_shape_text(shape))
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    note = slide.notes_slide.notes_text_frame.text.strip()
                    if note:
                        slide_texts.append(f"[Note: {note}]")
                if slide_texts:
                    text_runs.append(f"--- Slide {slide_num} ---\n" + "\n".join(slide_texts))
            res = "\n\n".join(text_runs).strip()
            if res:
                return res
        except Exception:
            pass

        # Try fallback text extractor for legacy .ppt
        fallback = _extract_ppt_fallback(raw_bytes)
        if fallback:
            return fallback

    return None


@app.post("/api/documents/upload")
async def upload_document(file: UploadFile = File(...)):
    raw = await file.read()
    text = _extract_text(file.filename, raw)
    if not text or not text.strip():
        raise HTTPException(status_code=400, detail="Could not extract text from document. Ensure file is a readable PPTX, PPT, PDF, DOCX, or TXT.")

    slide_count = text.count("--- Slide ")
    return {
        "filename": file.filename,
        "char_count": len(text),
        "slide_count": slide_count if slide_count > 0 else None,
        "text_preview": text[:500],
        "extracted_text": text,
    }


class DocFlashcardsModel(BaseModel):
    document_text: str
    n_cards: int = 12
    title: str = "Uploaded Document"


@app.post("/api/documents/flashcards")
def generate_document_flashcards(body: DocFlashcardsModel):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        res = llm_service.generate_flashcards(body.document_text, n_cards=body.n_cards)
        if not res["ok"] or not res.get("cards"):
            raise HTTPException(status_code=400, detail=res.get("error", "Failed to generate flashcards."))

        # Persist set
        fset = FlashcardSet(user_id=user.id, title=body.title, source_type="document")
        db.add(fset)
        db.flush()
        for c in res["cards"]:
            db.add(Flashcard(
                set_id=fset.id,
                user_id=user.id,
                front=c.get("front", ""),
                back=c.get("back", ""),
                difficulty=c.get("difficulty", "medium"),
            ))
        db.commit()

        return {
            "ok": True,
            "set_id": fset.id,
            "cards": res["cards"],
            "title": body.title,
            "provider": res.get("provider"),
            "provider_error": res.get("provider_error"),
        }
    finally:
        db.close()


class DocQuizModel(BaseModel):
    document_text: str
    n_questions: int = 8
    difficulty: str = "medium"
    title: str = "Document Quiz"


@app.post("/api/documents/quiz")
def generate_document_quiz(body: DocQuizModel):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        res = llm_service.generate_quiz_from_text(
            document_text=body.document_text,
            n_questions=body.n_questions,
            difficulty=body.difficulty,
        )
        if not res["ok"]:
            raise HTTPException(status_code=400, detail=res.get("error", "Failed to generate quiz."))

        quiz = Quiz(
            user_id=user.id,
            subject_id=None,
            topic_id=None,
            difficulty=body.difficulty,
            n_questions=len(res["questions"]),
            questions=json.dumps(res["questions"]),
            cache_key=None,
        )
        db.add(quiz)
        db.commit()
        db.refresh(quiz)

        return {
            "ok": True,
            "quiz_id": quiz.id,
            "title": body.title,
            "questions": res["questions"],
            "provider": res.get("provider"),
            "provider_error": res.get("provider_error"),
        }
    finally:
        db.close()


class RegisterDocSubjectModel(BaseModel):
    document_text: str
    subject_name: Optional[str] = None
    name: Optional[str] = None
    exam_date: str
    priority: int = 1


@app.post("/api/documents/register-subject")
def register_document_as_subject(body: RegisterDocSubjectModel):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        sname = (body.subject_name or body.name or "Uploaded Subject").strip()
        ext = llm_service.extract_topics_from_text(body.document_text, default_subject=sname)
        if not ext["ok"]:
            raise HTTPException(status_code=400, detail=ext.get("error", "Could not extract topics."))

        ok, msg = add_subject_with_topics(
            db, user.id,
            subject_name=sname,
            exam_date_str=body.exam_date,
            topics_list=ext["topics"],
            priority=body.priority,
        )
        if not ok:
            raise HTTPException(status_code=400, detail=msg)

        plan_res = apply_dynamic_timetable_update(db, user.id)
        return {
            "ok": True,
            "message": msg,
            "topics": ext["topics"],
            "plan_sessions_created": plan_res.get("created", 0),
            "provider": ext.get("provider"),
            "provider_error": ext.get("provider_error"),
        }
    finally:
        db.close()



@app.get("/api/flashcard-sets")
def get_flashcard_sets():
    db = get_db()
    try:
        user = _get_or_create_user(db)
        sets = (
            db.query(FlashcardSet)
            .filter(FlashcardSet.user_id == user.id)
            .order_by(FlashcardSet.created_at.desc())
            .all()
        )
        res = []
        for s in sets:
            cards = (
                db.query(Flashcard)
                .filter(Flashcard.set_id == s.id)
                .order_by(Flashcard.id)
                .all()
            )
            res.append({
                "id": s.id,
                "title": s.title,
                "created_at": s.created_at.strftime("%d %b %Y, %H:%M") if s.created_at else "—",
                "cards": [{"id": c.id, "front": c.front, "back": c.back, "difficulty": c.difficulty} for c in cards]
            })
        return res
    finally:
        db.close()


@app.delete("/api/flashcard-sets/{set_id}")
def delete_flashcard_set(set_id: int):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        fset = db.query(FlashcardSet).filter(FlashcardSet.id == set_id, FlashcardSet.user_id == user.id).first()
        if not fset:
            raise HTTPException(status_code=404, detail="Set not found.")
        db.delete(fset)
        db.commit()
        return {"ok": True, "message": "Set deleted."}
    finally:
        db.close()


# ─────────────────────────────────────────────────────────────────────────────
# 7. Model Lab & Predictor
# ─────────────────────────────────────────────────────────────────────────────

class PredictInputModel(BaseModel):
    study_hours: float
    prev_exam_score: float


@app.get("/api/predictor/suggest")
def get_session_suggestion(previous_score: float = Query(..., ge=0, le=100)):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        sug = session_suggestion(previous_score, user.session_min or config.DEFAULT_SESSION_MINUTES)
        return {
            "model_available": sug["model_available"],
            "benchmark_hours": sug["benchmark_hours"],
            "session_hours": sug["session_hours"],
            "suggested_session_count": sug["suggested_session_count"],
            "projected_hours": sug["projected_hours"],
            "pass_probability": to_percent(sug["pass_probability"]),
            "warning": None if sug["model_available"] else "Pass estimates unavailable — the trained model did not load.",
        }
    finally:
        db.close()


@app.get("/api/predictor/readiness")
def get_readiness():
    db = get_db()
    try:
        user = _get_or_create_user(db)
        return _readiness_payload(db, user)
    finally:
        db.close()


@app.post("/api/model/predict")
def run_model_prediction(body: PredictInputModel):
    model = load_model()
    res = predict(model, body.study_hours, body.prev_exam_score)
    if not res.get("ok"):
        raise HTTPException(status_code=400, detail=res.get("error", "Prediction failed."))

    prob = res.get("pass_probability") or 0.0
    return {
        "pass_probability": round(prob * 100, 1),
        "prediction_label": res.get("outcome_label", "Pass"),
        "is_pass": res.get("predicted_class") == 1,
    }


@app.get("/api/model/meta")
def get_model_metadata():
    meta = config.load_model_metadata()
    return meta


class VisualiserStateModel(BaseModel):
    study_hours: float = Field(..., ge=config.STUDY_HOURS_MIN, le=config.STUDY_HOURS_MAX)
    previous_score: float = Field(..., ge=config.EXAM_SCORE_MIN, le=config.EXAM_SCORE_MAX)


@app.get("/api/model-visualiser/overview")
def get_visualiser_overview():
    return visualiser_overview()


@app.post("/api/model-visualiser/state")
def get_visualiser_state(body: VisualiserStateModel):
    res = visualiser_state(body.study_hours, body.previous_score)
    if not res.get("ok"):
        raise HTTPException(status_code=400, detail=res.get("error", "Prediction failed."))
    return res


# ─────────────────────────────────────────────────────────────────────────────
# 8. AI Doubt Solver (Groq LPU Inference)
# ─────────────────────────────────────────────────────────────────────────────

class SolveDoubtInputModel(BaseModel):
    question: str
    subject_id: Optional[int] = None
    topic_id: Optional[int] = None
    depth: Optional[str] = "standard"
    api_key_override: Optional[str] = None


class SaveGroqConfigInputModel(BaseModel):
    api_key: str
    model: Optional[str] = "llama-3.3-70b-versatile"


@app.post("/api/doubts/solve")
def solve_academic_doubt(body: SolveDoubtInputModel):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        subject_name = None
        topic_name = None

        if body.subject_id:
            subj = db.query(Subject).filter(Subject.id == body.subject_id, Subject.user_id == user.id).first()
            if subj:
                subject_name = subj.name

        if body.topic_id:
            top = db.query(Topic).filter(Topic.id == body.topic_id, Topic.user_id == user.id).first()
            if top:
                topic_name = top.name

        res = doubt_solver.solve_doubt(
            question=body.question,
            subject_name=subject_name,
            topic_name=topic_name,
            depth=body.depth or "standard",
            api_key_override=body.api_key_override,
        )

        if not res.get("ok"):
            raise HTTPException(status_code=400, detail=res.get("error", "Failed to solve doubt."))

        # Persist doubt to history if answered
        doubt_id = None
        if not res.get("is_setup_required"):
            doubt_rec = Doubt(
                user_id=user.id,
                subject_id=body.subject_id,
                topic_id=body.topic_id,
                question=body.question.strip(),
                answer=res["answer"],
                key_takeaways=json.dumps(res.get("key_takeaways", [])),
                practice_question=json.dumps(res.get("practice_question")) if res.get("practice_question") else None,
                depth=body.depth or "standard",
                model_used=res.get("model_used", "llama-3.3-70b-versatile"),
            )
            db.add(doubt_rec)
            db.commit()
            db.refresh(doubt_rec)
            doubt_id = doubt_rec.id

        return {
            "ok": True,
            "id": doubt_id,
            "answer": res["answer"],
            "key_takeaways": res.get("key_takeaways", []),
            "practice_question": res.get("practice_question"),
            "model_used": res.get("model_used"),
            "provider": res.get("provider"),
            "latency_ms": res.get("latency_ms", 0),
            "is_setup_required": res.get("is_setup_required", False),
        }
    finally:
        db.close()


@app.get("/api/doubts")
def get_doubt_history():
    db = get_db()
    try:
        user = _get_or_create_user(db)
        records = (
            db.query(Doubt)
            .filter(Doubt.user_id == user.id)
            .order_by(Doubt.created_at.desc())
            .limit(40)
            .all()
        )
        results = []
        for d in records:
            takeaways = []
            if d.key_takeaways:
                try:
                    takeaways = json.loads(d.key_takeaways)
                except Exception:
                    takeaways = [d.key_takeaways]

            practice_q = None
            if d.practice_question:
                try:
                    practice_q = json.loads(d.practice_question)
                except Exception:
                    practice_q = None

            results.append({
                "id": d.id,
                "question": d.question,
                "answer": d.answer,
                "key_takeaways": takeaways,
                "practice_question": practice_q,
                "depth": d.depth,
                "model_used": d.model_used,
                "subject_name": d.subject.name if d.subject else None,
                "topic_name": d.topic.name if d.topic else None,
                "created_at": d.created_at.isoformat() if d.created_at else None,
            })
        return results
    finally:
        db.close()


@app.delete("/api/doubts/{doubt_id}")
def delete_doubt_entry(doubt_id: int):
    db = get_db()
    try:
        user = _get_or_create_user(db)
        d = db.query(Doubt).filter(Doubt.id == doubt_id, Doubt.user_id == user.id).first()
        if not d:
            raise HTTPException(status_code=404, detail="Doubt entry not found.")
        db.delete(d)
        db.commit()
        return {"ok": True, "message": "Doubt removed from history."}
    finally:
        db.close()


@app.post("/api/doubts/config")
def update_groq_configuration(body: SaveGroqConfigInputModel):
    res = doubt_solver.save_groq_api_key_to_env(body.api_key, body.model or "llama-3.3-70b-versatile")
    if not res.get("ok"):
        raise HTTPException(status_code=500, detail=res.get("error", "Failed to update configuration."))
    return res


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)
